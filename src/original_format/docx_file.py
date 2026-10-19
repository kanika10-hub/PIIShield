"""
Word: the redacted copy of a DOCX file.

The file is read again by the DOCX extractor in locate mode
(text_sources.SourceMap), so each block's text maps back, character by
character, to the w:t elements it came from. A file whose text is not
the extracted document's is refused. Each finding is replaced there by
its tag, at the scanner's offsets: the tag takes the formatting of the
run the value starts in. Then:

    text box copies     Word keeps a second copy of each text box for old
                        readers (mc:Fallback): the same edits are made there
    comment authors     and their initials, and file properties
    images              the OCR words of each value are filled (images.py)
    charts, SmartArt    their values replaced in every part that holds
                        them (the chart's workbook, the diagram's drawing)
    every part          searched for every value the scan found (plan.py,
                        ooxml.sweep_tree): tracked deletions, field codes,
                        hyperlink targets, alt text, custom XML, embedded
                        packages
    SVG copies          of pictures removed: the redacted pixels show
    thumbnail           the preview picture of the first page is removed

    redact_docx(source_path, plan, output_path) -> (summary, errors)
"""

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from src.extractors.docx_extractor import TEXT_BOX_PARAGRAPHS, docx_blocks
from src.extractors.text_sources import (
    ImageSource, PartsSource, SourceMap, TextSource, docx_text_pieces, w
)
from src.original_format.ooxml import (
    MC_NS, PartTrees, Unread, drop_svg_copies, finder, image_areas, located_pieces, replace_pieces,
    sweep_part, sweep_relationships
)
from src.original_format.images import redact_image
from src.original_format.plan import ValueMatcher, span_triples, sweep_values


def check_located(sources, plan, file_name):
    """
    Refuse a file whose text is not the extracted document's: the
    findings' offsets would point at the wrong characters.
    """

    differ = [
        block_id
        for block_id, source in sources.blocks.items()
        if isinstance(source, TextSource)
        and (plan.blocks.get(block_id, (None, {}))[1].get("text") or "") != source.text()
    ]

    if differ:
        raise ValueError(
            f"{file_name} does not match its extracted document: the text of {len(differ)} "
            f"blocks differs ({', '.join(differ[:5])})"
        )


def fallback_copies(source, used):
    """
    The copies Word keeps of a text box for old readers (VML, in the
    mc:Fallback beside it) that hold the same text. used: the copies
    already given, so each is edited once.
    """

    box = next(source.paragraphs[0].iterancestors(w("txbxContent")), None) if source.paragraphs else None
    choice = next(box.iterancestors(f"{{{MC_NS}}}AlternateContent"), None) if box is not None else None

    if choice is None:
        return []

    for fallback in choice.iterchildren(f"{{{MC_NS}}}Fallback"):
        for other in fallback.iter(w("txbxContent")):
            copy = TextSource(TEXT_BOX_PARAGRAPHS(other), docx_text_pieces)

            if other not in used and copy.text() == source.text():
                used.add(other)
                return [copy]

    return []


def set_author(comment, author):
    """
    A comment's author, and its initials, which would give it away.
    """

    comment.set(qn("w:author"), author)

    if comment.get(qn("w:initials")) is not None:
        comment.set(qn("w:initials"), "")


def redact_docx(source_path, plan, output_path):
    """
    Write the redacted copy of the DOCX file the plan's document was
    extracted from. Returns (summary, errors): the counts of what was
    replaced, and each finding that could not be (block_id, finding_id,
    reason, never the value). Raises ValueError if the file does not
    match the document.
    """

    doc = Document(str(source_path))
    sources = SourceMap()
    docx_blocks(doc, Unread(), sources)
    check_located(sources, plan, source_path.name)

    trees = PartTrees(sources.trees)
    part_values = {}                # partname -> value -> tag, for charts and diagrams
    image_parts = {}                # partname -> (part, areas, block_ids)
    used_copies = set()
    errors = []
    summary = {"spans_replaced": 0, "images_redacted": 0, "values_swept": 0, "svg_copies_removed": 0}

    def failed(block_id, spans, reason):
        errors.extend(
            {"block_id": block_id, "finding_id": item.get("finding_id"), "reason": reason}
            for span in spans for item in span["items"]
        )

    for block_id, spans in plan.spans.items():
        source = sources.blocks.get(block_id)

        if block_id.startswith("property_"):
            setattr(doc.core_properties, block_id[len("property_"):], plan.replaced_slots[block_id])
            summary["spans_replaced"] += len(spans)

        elif block_id in plan.replaced_slots and source is not None:
            set_author(source, plan.replaced_slots[block_id])
            summary["spans_replaced"] += len(spans)

        elif isinstance(source, TextSource):
            try:
                for target in (source, *fallback_copies(source, used_copies)):
                    replace_pieces(located_pieces(target), span_triples(spans))
            except ValueError as error:
                failed(block_id, spans, str(error))
                continue

            summary["spans_replaced"] += len(spans)

        elif isinstance(source, PartsSource):
            text = plan.blocks[block_id][1].get("text") or ""

            for part in source.parts:
                values = part_values.setdefault(str(part.partname), {})

                for start, end, tag in span_triples(spans):
                    values.setdefault(text[start:end], tag)

            summary["spans_replaced"] += len(spans)

        elif isinstance(source, ImageSource):
            areas = image_areas(plan.blocks[block_id][1], spans)

            if areas is None:
                failed(block_id, spans, "its OCR words do not match its text, so the values cannot be placed")
                continue

            _, all_areas, block_ids = image_parts.setdefault(str(source.part.partname), (source.part, [], []))
            all_areas.extend(areas)
            block_ids.append((block_id, spans))

        else:
            failed(block_id, spans, "the block was not found in the file")

    for part, areas, block_ids in image_parts.values():
        try:
            part._blob = redact_image(part.blob, areas)
            summary["images_redacted"] += 1
            summary["spans_replaced"] += sum(len(spans) for _, spans in block_ids)
        except ValueError as error:
            for block_id, spans in block_ids:
                failed(block_id, spans, str(error))

    # Every value the scan found, wherever else it is in the file
    base_values = sweep_values(plan.report["findings"])
    package = doc.part.package

    for part in package.iter_parts():
        extra = part_values.get(str(part.partname))
        matcher = ValueMatcher({**base_values, **extra}) if extra else plan.values
        find = finder(matcher, plan)

        if matcher:
            summary["values_swept"] += sweep_part(part, trees, find)
            summary["values_swept"] += sweep_relationships(part.rels, find)

    if plan.values:
        summary["values_swept"] += sweep_relationships(package.rels, finder(plan.values, plan))

    # Office would show a picture's SVG copy, not its redacted pixels
    for part in list(package.iter_parts()):
        summary["svg_copies_removed"] += drop_svg_copies(part, trees)

    trees.write_back()

    # A picture of the first page, as it was
    for rel_id in [rel_id for rel_id, rel in package.rels.items() if rel.reltype == RT.THUMBNAIL]:
        del package.rels[rel_id]

    doc.save(str(output_path))

    return summary, errors
