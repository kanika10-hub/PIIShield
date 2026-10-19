
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import nsmap, qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from src.extractors.common import office_properties, page_text_metadata, text_method
from src.extractors.text_sources import (
    ImageSource, PartsSource, Sources, TextSource, docx_comment_pieces, docx_paragraph_pieces,
    docx_text_pieces, pieces_text
)
from src.models.document import (
    ExtractedDocument,
    DocumentPage,
    DocumentBlock
)
from src.ocr.image_ocr import ImageReader

# Plain lxml elements (e.g. w:sdt) need the namespace map passed to xpath()
NAMESPACES = {
    "w": nsmap["w"],
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
}

# Charts and SmartArt diagrams: a drawing whose graphic data names a part
# that holds the text. Word keeps a picture of each in mc:Fallback.
CHART_URI = NAMESPACES["c"]
DIAGRAM_URI = NAMESPACES["dgm"]
GRAPHIC_DATA = etree.XPath(".//a:graphicData[not(ancestor::mc:Fallback)]", namespaces=NAMESPACES)
GRAPHIC_PARTS = {
    CHART_URI: ("chart", etree.XPath("./c:chart/@r:id", namespaces=NAMESPACES)),
    DIAGRAM_URI: ("diagram", etree.XPath("./dgm:relIds/@r:dm", namespaces=NAMESPACES)),
}

# A chart's visible text: titles, series names and category labels. Its
# values come from an embedded workbook that is not read.
CHART_TEXT = etree.XPath(
    ".//a:t/text() | .//c:tx//c:v/text() | .//c:cat//c:v/text()", namespaces=NAMESPACES
)
# A SmartArt diagram's text: each node's paragraphs in the data part
DIAGRAM_PARAGRAPHS = etree.XPath(".//a:p", namespaces=NAMESPACES)
PARAGRAPH_TEXT = etree.XPath(".//a:t/text()", namespaces=NAMESPACES)
# The drawing of a diagram, which Word shows, is linked from the story
DIAGRAM_DRAWING_IDS = etree.XPath(".//*[local-name()='dataModelExt']/@relId")

# An embedded object (a spreadsheet, another document): never opened
OLE_OBJECTS = etree.XPath(".//w:object[not(ancestor::mc:Fallback)]", namespaces=NAMESPACES)
OLE_PROG_ID = etree.XPath(".//*[local-name()='OLEObject']/@ProgID")

# Word stores each text box and drawing twice: as a drawing, and as a
# VML copy inside mc:Fallback for old readers. Only the first is read.
TEXT_BOXES = etree.XPath(
    ".//w:txbxContent[not(ancestor::mc:Fallback)]", namespaces=NAMESPACES
)
BLIPS = etree.XPath(
    ".//a:blip[not(ancestor::mc:Fallback)][@r:embed]", namespaces=NAMESPACES
)
TEXT_BOX_PARAGRAPHS = etree.XPath(".//w:p", namespaces=NAMESPACES)

# A cell's paragraphs, including those of tables nested in it
CELL_PARAGRAPHS = etree.XPath(
    ".//w:p[not(ancestor::w:txbxContent)]", namespaces=NAMESPACES
)

LIST_LEVEL = etree.XPath("./w:pPr/w:numPr/w:ilvl/@w:val", namespaces=NAMESPACES)
IS_LIST = etree.XPath("boolean(./w:pPr/w:numPr)", namespaces=NAMESPACES)
ENDS_SECTION = etree.XPath("boolean(./w:pPr/w:sectPr)", namespaces=NAMESPACES)

# Footnotes and endnotes that are separators, not notes
NOTE_SEPARATORS = ("separator", "continuationSeparator", "continuationNotice")

# Where comments and notes are attached: the ids they are marked with
# in the text they belong to
ANCHOR_IDS = {
    "comment": etree.XPath(
        ".//w:commentRangeStart/@w:id | .//w:commentReference/@w:id", namespaces=NAMESPACES
    ),
    "footnote": etree.XPath(".//w:footnoteReference/@w:id", namespaces=NAMESPACES),
    "endnote": etree.XPath(".//w:endnoteReference/@w:id", namespaces=NAMESPACES),
}


def record_anchors(anchors, element, block_id):
    """
    Note in anchors ((kind, id) -> block_id) the comments, footnotes and
    endnotes attached to the text of element, a block's XML. The first
    block a comment covers is the one it is about.
    """

    for kind, query in ANCHOR_IDS.items():
        for note_id in query(element):
            anchors.setdefault((kind, str(note_id)), block_id)

# Every header and footer a section can define
HEADER_FOOTER_PARTS = (
    "header", "footer",
    "first_page_header", "first_page_footer",
    "even_page_header", "even_page_footer"
)


def content_control_name(sdt):
    """
    Return a readable name for a content control (w:sdt),
    e.g. "Table of Contents" or a cover page field alias.
    """

    for query in (
        "./w:sdtPr/w:docPartObj/w:docPartGallery/@w:val",
        "./w:sdtPr/w:alias/@w:val",
        "./w:sdtPr/w:tag/@w:val"
    ):
        values = sdt.xpath(query, namespaces=NAMESPACES)

        if values:
            return values[0]

    return "content_control"


def iter_block_elements(parent, content_control=None):
    """
    Yield (element, content_control) for every paragraph and table
    directly under parent, in document order.

    Block-level content controls such as the table of contents wrap
    their own paragraphs and tables, so their content is yielded in
    place, tagged with the control's name.
    """

    for element in parent.iterchildren():

        if element.tag in (qn("w:p"), qn("w:tbl")):
            yield element, content_control

        elif element.tag == qn("w:sdt"):
            for content in element.iterchildren(qn("w:sdtContent")):
                yield from iter_block_elements(
                    content,
                    content_control_name(element)
                )


def paragraph_text(paragraph):
    """
    Return the full text of a paragraph.

    paragraph.text skips runs inside inline content controls, which
    hold cover page fields such as the client name and project title.
    This reads every run in the paragraph instead, except runs inside
    text boxes: those are stored twice (drawing and VML fallback).
    See text_sources.docx_paragraph_pieces.
    """

    return pieces_text(docx_paragraph_pieces(paragraph._p))


def extract_table_cells(table, table_number, location="body", sources=Sources(), anchors=None):
    """
    Return one table_cell block per physical cell, in row order.

    python-docx repeats a merged cell for every grid position it
    covers. Each cell is emitted once, at its top-left position,
    with the number of rows and columns it spans. A table nested in a
    cell is part of that cell's text. anchors: see record_anchors.
    """

    spans = {}

    for row_number, row in enumerate(table.rows, start=1):
        for column_number, cell in enumerate(row.cells, start=1):

            span = spans.setdefault(cell._tc, {
                "cell": cell,
                "row": row_number,
                "column": column_number,
                "rows": set(),
                "columns": set()
            })

            span["rows"].add(row_number)
            span["columns"].add(column_number)

    cells = []

    for cell_order, span in enumerate(spans.values(), start=1):

        cell = span["cell"]
        block_id = f"table_{table_number}_row_{span['row']}_cell_{span['column']}"
        source = TextSource(CELL_PARAGRAPHS(cell._tc), docx_paragraph_pieces)
        sources.add(block_id, source)

        if anchors is not None:
            record_anchors(anchors, cell._tc, block_id)

        # Empty cells are kept so the table grid stays intact
        cells.append(
            DocumentBlock(
                block_id=block_id,
                block_type="table_cell",
                text=source.text(),
                order=cell_order,
                metadata={
                    "source": "native",
                    "location": location,
                    "row": span["row"],
                    "column": span["column"],
                    "row_span": len(span["rows"]),
                    "column_span": len(span["columns"])
                }
            )
        )

    return cells


def lines_source(paragraphs):
    """
    Text of paragraph elements read with lxml (text boxes, notes): one
    line per paragraph, without tabs or breaks.
    """

    return TextSource(paragraphs, docx_text_pieces)


class StoryReader:
    """
    Reads one story of the document (the body, or a header or footer)
    into blocks, in reading order. Each paragraph or table is followed
    by the text boxes, images, charts, diagrams and embedded objects
    inside it.
    """

    def __init__(self, story, prefix, metadata, reader, section=1, track_sections=False, sources=Sources(),
                 anchors=None):
        self.story = story                  # python-docx parent of its paragraphs
        self.prefix = prefix                # block ID prefix, "" for the body
        self.metadata = metadata            # location (and part) of every block
        self.reader = reader
        self.sources = sources              # where each block's text is in the file
        self.anchors = {} if anchors is None else anchors   # see record_anchors
        self.section = section
        self.track_sections = track_sections
        self.counts = {}
        self.blocks = []

    def next_number(self, kind):
        self.counts[kind] = self.counts.get(kind, 0) + 1
        return self.counts[kind]

    def block_metadata(self, content_control=None):

        metadata = {"source": "native", **self.metadata, "section": self.section}

        if content_control:
            metadata["content_control"] = content_control

        return metadata

    def read(self, element):

        for child, content_control in iter_block_elements(element):

            if child.tag == qn("w:p"):
                self.paragraph(child, content_control)
            else:
                self.table(child, content_control)

            self.text_boxes(child)
            self.images(child)
            self.graphics(child)

            # A paragraph holding section properties ends its section
            if self.track_sections and child.tag == qn("w:p") and ENDS_SECTION(child):
                self.section += 1

        return self.blocks

    def paragraph(self, element, content_control):

        paragraph = Paragraph(element, self.story)
        source = TextSource([element], docx_paragraph_pieces)
        text = source.text()

        if not text:
            return

        try:
            style = paragraph.style.name if paragraph.style is not None else ""
        except (KeyError, ValueError, AttributeError):
            style = ""

        if style.startswith("Heading") or style == "Title":
            block_type = "heading"
        elif IS_LIST(element) or style.startswith("List"):
            block_type = "list_item"
        else:
            block_type = "paragraph"

        number = self.next_number("paragraph")
        metadata = {**self.block_metadata(content_control), "style": style, "paragraph_index": number}
        self.sources.add(f"{self.prefix}paragraph_{number}", source)
        record_anchors(self.anchors, element, f"{self.prefix}paragraph_{number}")

        if block_type == "list_item":
            # Set on the paragraph, or implied by the style ("List Bullet 2")
            levels = LIST_LEVEL(element)
            style_level = style.rsplit(" ", 1)[-1]

            if levels:
                metadata["level"] = int(levels[0])
            elif style_level.isdigit():
                metadata["level"] = int(style_level) - 1
            else:
                metadata["level"] = 0

        self.blocks.append(DocumentBlock(
            block_id=f"{self.prefix}paragraph_{number}",
            block_type=block_type,
            text=text,
            metadata=metadata
        ))

    def table(self, element, content_control):

        table = Table(element, self.story)
        number = self.next_number("table")

        self.blocks.append(DocumentBlock(
            block_id=f"{self.prefix}table_{number}",
            block_type="table",
            children=extract_table_cells(
                table, f"{self.prefix}{number}" if self.prefix else number,
                self.metadata["location"], self.sources, self.anchors
            ),
            metadata={
                **self.block_metadata(content_control),
                "rows": len(table.rows),
                "columns": len(table.columns)
            }
        ))

    def text_boxes(self, element):

        for box in TEXT_BOXES(element):
            source = lines_source(TEXT_BOX_PARAGRAPHS(box))
            text = source.text()

            if text:
                block_id = f"{self.prefix}textbox_{self.next_number('textbox')}"
                self.sources.add(block_id, source)
                record_anchors(self.anchors, box, block_id)

                self.blocks.append(DocumentBlock(
                    block_id=block_id,
                    block_type="paragraph",
                    text=text,
                    metadata={**self.block_metadata(), "location": "text_box", "in": self.metadata["location"]}
                ))

    def images(self, element):

        for blip in BLIPS(element):
            part = self.story.part.related_parts.get(blip.get(qn("r:embed")))

            if part is None:
                continue                    # a linked image: not in the file

            metadata = {**self.block_metadata(), "image_part": str(part.partname)}
            del metadata["source"]

            block_id = f"{self.prefix}image_{self.next_number('image')}"
            self.sources.add(block_id, ImageSource(part, [blip]))

            self.blocks.append(self.reader.block(
                block_id, str(part.partname), part.blob, metadata=metadata
            ))

    def graphics(self, element):
        """
        Charts and SmartArt diagrams, with the text they show, and embedded
        objects, which are reported but not opened.
        """

        for data in GRAPHIC_DATA(element):
            kind, relationship_ids = GRAPHIC_PARTS.get(data.get("uri"), (None, None))
            ids = relationship_ids(data) if kind else []
            part = self.story.part.related_parts.get(ids[0]) if ids else None

            if part is None:
                continue

            xml = etree.fromstring(part.blob)
            lines = (
                CHART_TEXT(xml) if kind == "chart"
                else ["".join(PARAGRAPH_TEXT(paragraph)) for paragraph in DIAGRAM_PARAGRAPHS(xml)]
            )

            block_id = f"{self.prefix}{kind}_{self.next_number(kind)}"

            # Every part that holds the text: a chart's own parts (its
            # embedded workbook), and the drawing Word shows for a diagram
            if kind == "chart":
                held = [rel.target_part for rel in part.rels.values() if not rel.is_external]
            else:
                held = [
                    self.story.part.related_parts[drawing_id]
                    for drawing_id in DIAGRAM_DRAWING_IDS(xml)
                    if drawing_id in self.story.part.related_parts
                ]

            self.sources.add(block_id, PartsSource([part, *held]))

            self.blocks.append(DocumentBlock(
                block_id=block_id,
                block_type=kind,
                text="\n".join(dict.fromkeys(line.strip() for line in lines if line.strip())),
                metadata={**self.block_metadata(), "graphic_part": str(part.partname)}
            ))

        for ole in OLE_OBJECTS(element):
            prog_id = OLE_PROG_ID(ole)

            self.blocks.append(DocumentBlock(
                block_id=f"{self.prefix}object_{self.next_number('object')}",
                block_type="embedded_object",
                metadata={**self.block_metadata(), "prog_id": prog_id[0] if prog_id else None}
            ))


def header_footer_blocks(doc, reader, sources=Sources(), anchors=None):
    """
    Blocks from every section's headers and footers.

    A header or footer linked to the previous section repeats that
    section's content and is skipped. Ones Word does not display (a
    first-page header while "different first page" is off) are still
    read, since their text is in the file; metadata.visible says which.
    anchors: see record_anchors.
    """

    blocks = []
    odd_and_even = doc.settings.odd_and_even_pages_header_footer

    for section_number, section in enumerate(doc.sections, start=1):
        for name in HEADER_FOOTER_PARTS:
            part = getattr(section, name)

            if part.is_linked_to_previous:
                continue

            if name.startswith("first_page"):
                visible = section.different_first_page_header_footer
            elif name.startswith("even_page"):
                visible = odd_and_even
            else:
                visible = True

            story = StoryReader(
                part, f"section_{section_number}_{name}_",
                {
                    "location": "header" if name.endswith("header") else "footer",
                    "part": name,
                    "visible": visible
                },
                reader, section=section_number, sources=sources, anchors=anchors
            )

            blocks.extend(story.read(part._element))

    return blocks


def note_blocks(doc, relationship, kind, sources=Sources(), anchors=None):
    """
    One block per footnote or endnote, skipping the separators Word adds.
    metadata.anchor_block_id is the block whose text refers to the note.
    """

    try:
        part = doc.part.part_related_by(relationship)
    except KeyError:
        return []

    blocks = []

    for note in sources.xml(part):

        if note.get(qn("w:type")) in NOTE_SEPARATORS:
            continue

        source = lines_source(TEXT_BOX_PARAGRAPHS(note))
        text = source.text()

        if text:
            block_id = f"{kind}_{note.get(qn('w:id'))}"
            sources.add(block_id, source)
            metadata = {"source": "native", "location": kind, "note_id": note.get(qn("w:id"))}
            anchor = (anchors or {}).get((kind, str(note.get(qn("w:id")))))

            if anchor:
                metadata["anchor_block_id"] = anchor

            blocks.append(DocumentBlock(
                block_id=block_id,
                block_type="paragraph",
                text=text,
                metadata=metadata
            ))

    return blocks


def comment_blocks(doc, sources=Sources(), anchors=None):
    """
    One block per comment. The author is kept in metadata, and
    metadata.anchor_block_id is the block the comment is about.
    """

    # Reading doc.comments would add an empty comments part
    if not any(rel.reltype == RT.COMMENTS for rel in doc.part.rels.values()):
        return []

    blocks = []

    for comment in doc.comments:
        source = TextSource([paragraph._p for paragraph in comment.paragraphs], docx_comment_pieces)
        text = source.text()

        if text:
            block_id = f"comment_{comment.comment_id}"
            sources.add(block_id, source)
            # The comment element holds the author
            sources.add(f"{block_id}_author", comment._comment_elm)

            metadata = {
                "source": "native",
                "location": "comment",
                "comment_id": comment.comment_id,
                "author": comment.author,
                "initials": comment.initials
            }
            anchor = (anchors or {}).get(("comment", str(comment.comment_id)))

            if anchor:
                metadata["anchor_block_id"] = anchor

            blocks.append(DocumentBlock(
                block_id=block_id,
                block_type="paragraph",
                text=text,
                metadata=metadata
            ))

    return blocks


def docx_blocks(doc, reader, sources=Sources()):
    """
    Every block of an open document, in order. The original-format
    redactor calls this in locate mode, with a SourceMap.
    """

    # The body in reading order, then headers and footers, then notes and
    # comments: they sit outside the body's reading order, so each note and
    # comment names the block it belongs to (metadata.anchor_block_id)
    anchors = {}
    blocks = (
        StoryReader(
            doc, "", {"location": "body"}, reader, track_sections=True, sources=sources, anchors=anchors
        ).read(doc.element.body)
        + header_footer_blocks(doc, reader, sources, anchors)
    )
    blocks += (
        note_blocks(doc, RT.FOOTNOTES, "footnote", sources, anchors)
        + note_blocks(doc, RT.ENDNOTES, "endnote", sources, anchors)
        + comment_blocks(doc, sources, anchors)
    )

    for order, block in enumerate(blocks, start=1):
        block.order = order

    return blocks


def extract_docx(file_path):

    doc = Document(file_path)
    blocks = docx_blocks(doc, ImageReader())

    # DOCX has no fixed page boundaries: the whole document is page 1
    page = DocumentPage(
        page_number=1,
        blocks=blocks,
        metadata={
            "content_type": "docx_body",
            "extraction_method": text_method(blocks),
            **page_text_metadata(blocks)
        }
    )

    return ExtractedDocument(
        document_id=Path(file_path).stem,
        file_name=Path(file_path).name,
        source_type="DOCX",
        extraction_method=page.metadata["extraction_method"],
        pages=[page],
        metadata={
            "page_count": 1,
            "section_count": len(doc.sections),
            "properties": office_properties(doc.core_properties),
            "limitations": []
        }
    )
