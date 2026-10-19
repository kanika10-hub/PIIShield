from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

from src.extractors.common import combined_method, office_properties, page_text_metadata, text_method
from src.extractors.text_sources import (
    ImageSource, PartsSource, Sources, TextSource, pptx_comment_pieces, pptx_paragraph_pieces
)
from src.models.document import (
    ExtractedDocument,
    DocumentPage,
    DocumentBlock,
    BoundingBox
)
from src.ocr.image_ocr import ImageReader


DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DRAWING_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

# SmartArt text: each node's paragraphs in the diagram data part
DIAGRAM_PARAGRAPHS = etree.XPath(".//a:p", namespaces=DRAWING_NS)
PARAGRAPH_TEXT = etree.XPath(".//a:t/text()", namespaces=DRAWING_NS)

# The drawing PowerPoint shows for a SmartArt diagram: a copy of its text,
# named by relId in the data part and related from the slide
DIAGRAM_DRAWING_IDS = etree.XPath(".//*[local-name()='dataModelExt']/@relId")
BLIPS = etree.XPath(".//a:blip", namespaces=DRAWING_NS)
EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

# Comments: PowerPoint's classic ones (p:cm, authors in p:cmAuthor) and
# the threaded ones of PowerPoint 365 (p188:cm with p188:reply, authors
# in p188:author)
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
P188_NS = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
THREADED_COMMENTS = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
THREADED_AUTHORS = "http://schemas.microsoft.com/office/2018/10/relationships/authors"
COMMENT_PARAGRAPHS = etree.XPath("p188:txBody/a:p", namespaces={"p188": P188_NS, **DRAWING_NS})

# Placeholders whose paragraphs are bullets unless they say otherwise
BULLET_PLACEHOLDERS = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)

# A shape inside a group has its position in the group's own coordinate
# space. A transform (offset_x, offset_y, scale_x, scale_y) maps it to
# the slide: slide = offset + position * scale, in EMU.
SLIDE_TRANSFORM = (0, 0, 1, 1)


def get_bounding_box(shape, transform=SLIDE_TRANSFORM):
    """
    Return the shape position on the slide in points, or None if
    unavailable.
    """

    if None in (shape.left, shape.top, shape.width, shape.height):
        return None

    offset_x, offset_y, scale_x, scale_y = transform

    return BoundingBox(
        x=Emu(int(offset_x + shape.left * scale_x)).pt,
        y=Emu(int(offset_y + shape.top * scale_y)).pt,
        width=Emu(int(shape.width * scale_x)).pt,
        height=Emu(int(shape.height * scale_y)).pt
    )


def group_transform(group, transform):
    """
    The transform for shapes inside a group shape.
    """

    xfrm = group._element.grpSpPr.xfrm

    if xfrm is None or xfrm.chOff is None or xfrm.chExt is None or not xfrm.chExt.cx or not xfrm.chExt.cy:
        return transform

    offset_x, offset_y, scale_x, scale_y = transform
    group_scale_x = xfrm.ext.cx / xfrm.chExt.cx
    group_scale_y = xfrm.ext.cy / xfrm.chExt.cy

    return (
        offset_x + (xfrm.off.x - xfrm.chOff.x * group_scale_x) * scale_x,
        offset_y + (xfrm.off.y - xfrm.chOff.y * group_scale_y) * scale_y,
        scale_x * group_scale_x,
        scale_y * group_scale_y
    )


def find_title_shape(slide):
    """
    Find the shape that acts as the slide title.

    Uses the title placeholder if there is one. Otherwise falls back
    to the first text shape with the largest font size on the slide.
    """

    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE
        ):
            return shape

    title_shape = None
    largest_size = 0

    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue

        sizes = [
            run.font.size.pt
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size
        ]

        if sizes and max(sizes) > largest_size:
            largest_size = max(sizes)
            title_shape = shape

    return title_shape


def shape_metadata(shape, location="body"):

    metadata = {
        "source": "native",
        "location": location,
        "shape_id": shape.shape_id,
        "shape_name": shape.name
    }

    if shape.is_placeholder:
        metadata["placeholder"] = str(shape.placeholder_format.type)

    return metadata


def clean_text(text):
    # python-pptx gives a line break inside a paragraph as a vertical tab
    return text.replace("\v", "\n")


def is_bullet(paragraph, shape):
    """
    True if the paragraph shows a bullet or number: set on the paragraph
    itself, or inherited by a body placeholder from the slide master.
    """

    properties = paragraph._p.pPr

    if properties is not None:
        for child in properties:
            name = etree.QName(child).localname

            if name in ("buChar", "buAutoNum", "buBlip"):
                return True

            if name == "buNone":
                return False

    return shape.is_placeholder and shape.placeholder_format.type in BULLET_PLACEHOLDERS


def text_blocks(shape, base_id, block_type, transform, location, sources=Sources()):
    """
    One block per non-empty paragraph of a shape's text frame.
    """

    blocks = []

    for number, paragraph in enumerate(shape.text_frame.paragraphs, start=1):
        # clean_text(paragraph.text).strip(), piece by piece
        source = TextSource([paragraph._p], pptx_paragraph_pieces)
        text = source.text()

        if not text:
            continue

        bullet = block_type == "paragraph" and is_bullet(paragraph, shape)

        metadata = {**shape_metadata(shape, location), "paragraph": number}

        if bullet:
            metadata["level"] = paragraph.level

        sources.add(f"{base_id}_p{number}", source)

        blocks.append(DocumentBlock(
            block_id=f"{base_id}_p{number}",
            block_type="list_item" if bullet else block_type,
            text=text,
            bounding_box=get_bounding_box(shape, transform),
            metadata=metadata
        ))

    return blocks


def table_block(shape, table_id, transform, sources=Sources()):
    """
    A table block with one table_cell per cell. A merged cell is one
    cell at its top-left position, with the rows and columns it spans.
    """

    table = shape.table
    cells = []

    for row_index, row in enumerate(table.rows, start=1):
        for column_index, cell in enumerate(row.cells, start=1):

            if cell.is_spanned:
                continue

            # clean_text(cell.text), read without adding a text body to an empty cell
            body = cell._tc.txBody
            source = TextSource(
                body.p_lst if body is not None else [], pptx_paragraph_pieces, strip=False
            )
            block_id = f"{table_id}_cell_{row_index}_{column_index}"
            sources.add(block_id, source)

            cells.append(DocumentBlock(
                block_id=block_id,
                block_type="table_cell",
                text=source.text(),
                order=len(cells) + 1,
                metadata={
                    "source": "native",
                    "location": "body",
                    "row": row_index,
                    "column": column_index,
                    "row_span": cell.span_height,
                    "column_span": cell.span_width
                }
            ))

    return DocumentBlock(
        block_id=table_id,
        block_type="table",
        bounding_box=get_bounding_box(shape, transform),
        children=cells,
        metadata={
            **shape_metadata(shape),
            "rows": len(table.rows),
            "columns": len(table.columns)
        }
    )


def chart_block(shape, block_id, transform, sources=Sources()):
    """
    A chart's visible text: title, axis titles, series names and
    category labels. The values come from an embedded workbook that is
    not read.
    """

    chart = shape.chart
    lines = []

    # The chart part, and the parts it holds (the embedded workbook)
    chart_part = shape.chart_part
    sources.add(block_id, PartsSource([chart_part] + [
        rel.target_part for rel in chart_part.rels.values() if not rel.is_external
    ]))

    if chart.has_title and chart.chart_title.has_text_frame:
        lines.append(chart.chart_title.text_frame.text)

    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
            if axis.has_title:
                lines.append(axis.axis_title.text_frame.text)
        except (ValueError, AttributeError):
            pass                            # pie charts have no axes

    for plot in chart.plots:
        lines.extend(str(series.name) for series in plot.series if series.name)
        lines.extend(str(category) for category in plot.categories if category is not None)

    text = "\n".join(dict.fromkeys(clean_text(line).strip() for line in lines if str(line).strip()))

    return DocumentBlock(
        block_id=block_id,
        block_type="chart",
        text=text,
        bounding_box=get_bounding_box(shape, transform),
        metadata={**shape_metadata(shape), "chart_type": str(chart.chart_type)}
    )


def diagram_block(shape, block_id, transform, sources=Sources()):
    """
    A SmartArt diagram's text, from its data part: one line per paragraph.
    """

    relationship_ids = shape._element.xpath(".//dgm:relIds/@r:dm")
    text = ""

    if relationship_ids:
        data_part = shape.part.related_part(relationship_ids[0])
        data = etree.fromstring(data_part.blob)
        text = "\n".join(
            line
            for line in ("".join(PARAGRAPH_TEXT(paragraph)).strip() for paragraph in DIAGRAM_PARAGRAPHS(data))
            if line
        )

        # The data, and the drawing PowerPoint shows: both hold the text
        drawing_parts = [
            shape.part.related_part(drawing_id)
            for drawing_id in DIAGRAM_DRAWING_IDS(data)
            if drawing_id in shape.part.rels
        ]
        sources.add(block_id, PartsSource([data_part, *drawing_parts]))

    return DocumentBlock(
        block_id=block_id,
        block_type="diagram",
        text=text,
        bounding_box=get_bounding_box(shape, transform),
        metadata=shape_metadata(shape)
    )


def picture_block(shape, block_id, transform, reader, sources=Sources()):
    """
    An image_ocr block with the picture's OCR text (word boxes in image
    pixels; bounding_box is the picture on the slide), or an image
    block if it cannot be read.
    """

    metadata = {**shape_metadata(shape), "image_part": None}
    image_bytes = b""

    try:
        part = shape.part.related_part(shape._element.blip_rId)
        metadata["image_part"] = str(part.partname)
        image_bytes = part.blob
        sources.add(block_id, ImageSource(part, BLIPS(shape._element)))
    except (KeyError, AttributeError, ValueError):
        pass                                # a linked picture: not in the file

    del metadata["source"]

    return reader.block(
        block_id, metadata["image_part"] or block_id, image_bytes,
        bounding_box=get_bounding_box(shape, transform),
        metadata=metadata
    )


def shape_blocks(shapes, slide_number, title_shape, reader, transform=SLIDE_TRANSFORM, prefix="",
                 sources=Sources()):
    """
    Blocks for a slide's shapes, including shapes inside groups.
    Block IDs number the shapes in order; a shape inside a group is
    numbered within it ("3_2").
    """

    blocks = []

    for index, shape in enumerate(shapes, start=1):
        path = f"{prefix}{index}"
        uri = shape._element.xpath("./a:graphic/a:graphicData/@uri")

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            blocks.extend(shape_blocks(
                shape.shapes, slide_number, title_shape, reader,
                group_transform(shape, transform), f"{path}_", sources
            ))

        elif shape.has_text_frame and shape.text_frame.text.strip():
            blocks.extend(text_blocks(
                shape, f"slide{slide_number}_text{path}",
                "heading" if shape == title_shape else "paragraph",
                transform,
                "text_box" if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX else "body",
                sources
            ))

        elif getattr(shape, "has_table", False) and shape.has_table:
            blocks.append(table_block(shape, f"slide{slide_number}_table{path}", transform, sources))

        elif getattr(shape, "has_chart", False) and shape.has_chart:
            blocks.append(chart_block(shape, f"slide{slide_number}_chart{path}", transform, sources))

        elif uri and uri[0] == DIAGRAM_URI:
            blocks.append(diagram_block(shape, f"slide{slide_number}_diagram{path}", transform, sources))

        elif shape.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
            blocks.append(DocumentBlock(
                block_id=f"slide{slide_number}_object{path}",
                block_type="embedded_object",
                bounding_box=get_bounding_box(shape, transform),
                metadata={**shape_metadata(shape), "prog_id": shape.ole_format.prog_id}
            ))

        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
            shape.is_placeholder and hasattr(shape, "image")
        ):
            blocks.append(picture_block(
                shape, f"slide{slide_number}_image{path}", transform, reader, sources
            ))

    return blocks


def notes_blocks(slide, slide_number, sources=Sources()):
    """
    The slide's speaker notes, one block per paragraph.
    """

    if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
        return []

    blocks = []

    for number, paragraph in enumerate(slide.notes_slide.notes_text_frame.paragraphs, start=1):
        source = TextSource([paragraph._p], pptx_paragraph_pieces)
        text = source.text()

        if text:
            sources.add(f"slide{slide_number}_notes_p{number}", source)

            blocks.append(DocumentBlock(
                block_id=f"slide{slide_number}_notes_p{number}",
                block_type="paragraph",
                text=text,
                metadata={"source": "native", "location": "notes", "paragraph": number}
            ))

    return blocks


def template_image_blocks(slide, slide_number, reader, seen, sources=Sources()):
    """
    The images the slide shows that no picture block reads: its
    background and picture fills, and the pictures, fills and background
    its layout and master put on it (a logo, a backdrop). Each image is
    read once, on the first slide that shows it. seen: the image parts
    already read, updated here.
    """

    layout = slide.slide_layout
    blocks = []

    for kind, owner in (("fill", slide), ("layout", layout), ("master", layout.slide_master)):
        images = {}                         # partname -> (part, the a:blip elements that show it)

        for blip in BLIPS(owner._element):
            rel = owner.part.rels.get(blip.get(EMBED))

            if rel is not None and not rel.is_external and rel.reltype == RT.IMAGE:
                images.setdefault(str(rel.target_part.partname), (rel.target_part, []))[1].append(blip)

        for name, (part, blips) in images.items():
            if name in seen:
                continue

            seen.add(name)
            block_id = f"slide{slide_number}_{kind}_image{len(blocks) + 1}"
            sources.add(block_id, ImageSource(part, blips))

            blocks.append(reader.block(
                block_id, name, part.blob,
                metadata={"location": "body" if kind == "fill" else "template", "image_part": name}
            ))

    return blocks


def comment_authors(slide, sources=Sources()):
    """
    (kind, author id) -> the author's element, for the presentation's
    classic and threaded comment authors.
    """

    authors = {}

    for rel in slide.part.package.presentation_part.rels.values():
        if rel.is_external or rel.reltype not in (RT.COMMENT_AUTHORS, THREADED_AUTHORS):
            continue

        kind = "classic" if rel.reltype == RT.COMMENT_AUTHORS else "threaded"

        for author in sources.xml(rel.target_part).iter(f"{{{P_NS}}}cmAuthor", f"{{{P188_NS}}}author"):
            authors[(kind, author.get("id"))] = author

    return authors


def comment_blocks(slide, slide_number, sources=Sources()):
    """
    One block per comment on the slide, replies included. The author's
    name is kept in metadata, as for Word comments.
    """

    # By part name: the order a saved copy reads them in too
    rels = sorted(
        (
            rel for rel in slide.part.rels.values()
            if not rel.is_external and rel.reltype in (RT.COMMENTS, THREADED_COMMENTS)
        ),
        key=lambda rel: str(rel.target_part.partname)
    )

    if not rels:
        return []

    authors = comment_authors(slide, sources)
    blocks = []

    for rel in rels:
        root = sources.xml(rel.target_part)

        if rel.reltype == RT.COMMENTS:
            entries = [
                ("classic", comment, TextSource([text], pptx_comment_pieces))
                for comment in root.iter(f"{{{P_NS}}}cm")
                for text in comment.iterchildren(f"{{{P_NS}}}text")
            ]
        else:
            entries = [
                ("threaded", comment, TextSource(COMMENT_PARAGRAPHS(comment), pptx_paragraph_pieces))
                for comment in root.iter(f"{{{P188_NS}}}cm", f"{{{P188_NS}}}reply")
            ]

        for kind, comment, source in entries:
            text = source.text()

            if not text:
                continue

            block_id = f"slide{slide_number}_comment{len(blocks) + 1}"
            sources.add(block_id, source)

            metadata = {
                "source": "native",
                "location": "comment",
                "comment_id": comment.get("idx") or comment.get("id"),
            }
            author = authors.get((kind, comment.get("authorId")))

            if author is not None:
                # The author's element holds the name
                sources.add(f"{block_id}_author", author)
                metadata["author"] = author.get("name") or ""
                metadata["initials"] = author.get("initials")

            blocks.append(DocumentBlock(
                block_id=block_id,
                block_type="paragraph",
                text=text,
                metadata=metadata
            ))

    return blocks


def slide_blocks(slide, slide_number, reader, sources=Sources(), seen=None):
    """
    The blocks of one slide, in reading order, then the images its
    template puts on it, its speaker notes and its comments. seen: the
    image parts earlier slides read, updated here. The original-format
    redactor calls this in locate mode, with a SourceMap.
    """

    seen = set() if seen is None else seen

    blocks = shape_blocks(
        slide.shapes, slide_number, find_title_shape(slide), reader, sources=sources
    )

    # Reading order: top to bottom, then left to right. Paragraphs of
    # one shape share its position and keep their order.
    blocks.sort(key=lambda block: (
        (round(block.bounding_box.y), block.bounding_box.x)
        if block.bounding_box else (float("inf"), 0)
    ))

    seen.update(block.metadata["image_part"] for block in blocks if block.metadata.get("image_part"))
    blocks.extend(template_image_blocks(slide, slide_number, reader, seen, sources))
    blocks.extend(notes_blocks(slide, slide_number, sources))
    blocks.extend(comment_blocks(slide, slide_number, sources))

    for order, block in enumerate(blocks, start=1):
        block.order = order

    return blocks


def extract_pptx(file_path: str) -> ExtractedDocument:

    file_path = Path(file_path)

    presentation = Presentation(file_path)
    reader = ImageReader()
    seen = set()
    pages = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        blocks = slide_blocks(slide, slide_number, reader, seen=seen)

        pages.append(
            DocumentPage(
                page_number=slide_number,
                blocks=blocks,
                metadata={
                    "content_type": "slide",
                    "extraction_method": text_method(blocks),
                    "width": Emu(presentation.slide_width).pt,
                    "height": Emu(presentation.slide_height).pt,
                    "layout": slide.slide_layout.name,
                    "hidden": slide._element.get("show") == "0",
                    **page_text_metadata(blocks)
                }
            )
        )

    return ExtractedDocument(
        document_id=file_path.stem,
        file_name=file_path.name,
        source_type="PPTX",
        extraction_method=combined_method(
            page.metadata["extraction_method"] for page in pages
        ) if pages else "native",
        pages=pages,
        metadata={
            "page_count": len(pages),
            "properties": office_properties(presentation.core_properties),
            "limitations": []
        }
    )
